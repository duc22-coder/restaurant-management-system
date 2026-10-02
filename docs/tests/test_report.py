"""Kiểm tra nguồn báo cáo và cấu trúc DOCX, không cần khởi chạy ứng dụng."""
from __future__ import annotations

import contextlib
import io
from pathlib import Path
import re
import struct
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
import zipfile
import zlib

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm

DOCS = Path(__file__).resolve().parents[1]
REPO = DOCS.parent
sys.path.insert(0, str(DOCS))
import md2docx

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}


def fixture_png(width=100, height=50):
    def chunk(kind, data):
        payload = kind + data
        return struct.pack(">I", len(data)) + payload + struct.pack(">I", zlib.crc32(payload))
    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    pixels = (b"\0" + b"\xff\xff\xff" * width) * height
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header)
            + chunk(b"IDAT", zlib.compress(pixels)) + chunk(b"IEND", b""))


class ConverterTests(unittest.TestCase):
    def test_inline_formats_and_line_breaks(self):
        document = md2docx.build_document("Nội dung **in đậm**, *in nghiêng*, `Order`.<br>Dòng 2")
        paragraph = document.paragraphs[0]
        self.assertEqual(paragraph.text, "Nội dung in đậm, in nghiêng, Order.\nDòng 2")
        self.assertTrue(next(r for r in paragraph.runs if r.text == "in đậm").bold)
        self.assertTrue(next(r for r in paragraph.runs if r.text == "in nghiêng").italic)
        self.assertEqual(next(r for r in paragraph.runs if r.text == "Order").font.name,
                         "DejaVu Sans Mono")

    def test_tables_keep_first_body_row_and_alignment(self):
        markdown = "| Trái | Giữa | Phải |\n| :--- | :---: | ---: |\n| A | B<br>C | D |"
        document = md2docx.build_document(markdown)
        table = document.tables[0]
        self.assertEqual(len(table.rows), 2)
        self.assertEqual(table.cell(1, 1).text, "B\nC")
        self.assertEqual(table.cell(1, 1).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(table.cell(1, 2).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        root = ET.fromstring(document.element.xml)
        self.assertEqual(len(root.findall(".//w:tblHeader", NS)), 1)
        self.assertEqual(len(root.findall(".//w:cantSplit", NS)), 2)
        self.assertFalse(root.findall(".//w:pPr/w:pPr", NS))
        self.assertFalse(root.findall(".//w:r/w:b", NS))

    def test_malformed_table_fails(self):
        with self.assertRaises(ValueError):
            md2docx.build_document("| A | B |\n| --- | --- |\n| Thiếu ô |")

    def test_page_orientation_and_size(self):
        markdown = "# Dọc\n<!-- landscape -->\n# Ngang\n<!-- portrait -->\n# Dọc lại"
        document = md2docx.build_document(markdown)
        self.assertEqual([s.orientation for s in document.sections],
                         [WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT])
        self.assertAlmostEqual(document.sections[1].page_width, Mm(297), delta=635)
        self.assertAlmostEqual(document.sections[2].page_width, Mm(210), delta=635)

    def test_initial_landscape_does_not_make_blank_section(self):
        document = md2docx.build_document("<!-- landscape -->\n# Trang đầu ngang")
        self.assertEqual(len(document.sections), 1)
        self.assertEqual(document.sections[0].orientation, WD_ORIENT.LANDSCAPE)

    def test_consecutive_and_trailing_breaks_do_not_add_blank_pages(self):
        document = md2docx.build_document(
            "# Đầu\n<!-- pagebreak -->\n<!-- landscape -->\n# Sau\n<!-- pagebreak -->")
        self.assertEqual(len(document.sections), 2)
        root = ET.fromstring(document.element.xml)
        self.assertFalse(root.findall('.//w:br[@w:type="page"]', NS))

    def test_page_break_is_on_heading_not_an_empty_paragraph(self):
        document = md2docx.build_document("# Đầu\n<!-- pagebreak -->\n# Trang kế")
        self.assertEqual(len(document.paragraphs), 2)
        self.assertTrue(document.paragraphs[1].paragraph_format.page_break_before)
        root = ET.fromstring(document.element.xml)
        self.assertFalse(root.findall('.//w:br[@w:type="page"]', NS))

    def test_png_validation(self):
        self.assertEqual(md2docx.png_size(fixture_png()), (100, 50))
        for invalid in [b"not a PNG", b"\x89PNG\r\n\x1a\n"]:
            with self.assertRaises(ValueError):
                md2docx.png_size(invalid)

    def test_image_has_correct_ratio_alt_text_and_single_caption(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory)
            (base / "image.png").write_bytes(fixture_png())
            document = md2docx.build_document("![Hình thử](image.png)", base)
            picture = document.inline_shapes[0]
            self.assertAlmostEqual(picture.width / picture.height, 2, places=4)
            self.assertEqual(picture._inline.docPr.get("descr"), "Hình thử")
            self.assertEqual([p.text for p in document.paragraphs].count("Hình thử"), 1)
            self.assertTrue(document.paragraphs[0].paragraph_format.keep_with_next)
            self.assertEqual(document.paragraphs[1].style.name, "Caption")

    def test_missing_image_does_not_overwrite_existing_output(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.md", Path(directory) / "saved.docx"
            source.write_text("![Ảnh](missing.png)")
            output.write_bytes(b"existing document")
            with self.assertRaises(FileNotFoundError):
                md2docx.convert(source, output)
            self.assertEqual(output.read_bytes(), b"existing document")

    def test_successful_build_is_readable_docx(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / "source.md", Path(directory) / "saved.docx"
            source.write_text("# Tiếng Việt\nVăn bản có dấu.", encoding="utf-8")
            with contextlib.redirect_stdout(io.StringIO()):
                md2docx.convert(source, output)
            self.assertEqual(Document(output).paragraphs[0].text, "Tiếng Việt")


class ReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.markdown = md2docx.SRC.read_text(encoding="utf-8")
        cls.images = re.findall(r"^!\[(.*?)\]\((.*?)\)$", cls.markdown, re.M)
        cls.document = Document(md2docx.OUT)
        with zipfile.ZipFile(md2docx.OUT) as package:
            cls.parts = {name: package.read(name) for name in package.namelist()}
        cls.xml = ET.fromstring(cls.parts["word/document.xml"])

    def test_figure_numbers_are_unique_and_belong_to_chapter_three(self):
        numbers = [re.match(r"Hình 3\.(\d+) – ", caption).group(1)
                   for caption, _ in self.images]
        self.assertEqual(numbers, [str(n) for n in range(1, 14)])
        self.assertNotIn("Hình 2.", self.markdown)
        captions = [p.text for p in self.document.paragraphs if p.style.name == "Caption"]
        self.assertEqual(captions, [caption for caption, _ in self.images])

    def test_all_images_have_editable_sources_and_no_unused_assets(self):
        references = {Path(path).stem for _, path in self.images}
        sources = {p.stem for p in (DOCS / "bao-cao/uml").glob("*.puml")
                   if not p.name.startswith("_")}
        assets = {p.stem for p in (DOCS / "bao-cao/images").glob("*.png")}
        self.assertEqual(references, sources)
        self.assertEqual(references, assets)
        for _, path in self.images:
            self.assertGreater(min(md2docx.png_size((DOCS / path).read_bytes())), 100)

    def test_docx_embeds_updated_source_images_not_stale_ones(self):
        media = [data for name, data in self.parts.items() if name.startswith("word/media/")]
        self.assertEqual(len(media), 13)
        self.assertEqual(len(self.document.inline_shapes), 13)
        self.assertEqual(set(media), {(DOCS / path).read_bytes() for _, path in self.images})

    def test_word_xml_is_well_formed_and_paragraph_properties_not_nested(self):
        for name, data in self.parts.items():
            if name.endswith((".xml", ".rels")):
                ET.fromstring(data)
        self.assertFalse(self.xml.findall(".//w:pPr/w:pPr", NS))
        for tag in ("b", "i", "rFonts", "sz", "shd"):
            self.assertFalse(self.xml.findall(f".//w:r/w:{tag}", NS), tag)
        self.assertEqual(len(self.xml.findall(".//wp:docPr", NS)), 13)

    def test_report_is_free_of_placeholders_and_trial_watermarks(self):
        text = " ".join(self.xml.itertext())
        for forbidden in ("thiếu ảnh", "Evaluation Warning", "Spire.Doc", "Hết nội dung",
                          "Hình 2.", "class-tong-quan", "isAvailable()", "calculateTotal()"):
            self.assertNotIn(forbidden, text)
        self.assertEqual(len(self.document.tables), 4)
        self.assertEqual([s.orientation for s in self.document.sections],
                         [WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT])

    def test_class_diagrams_cover_real_entities_without_invented_methods(self):
        text = "\n".join(p.read_text() for p in (DOCS / "bao-cao/uml").glob("class-*.puml"))
        names = set(re.findall(r"^class (\w+)", text, re.M))
        entities = {p.stem for p in (REPO / "backend/src/main/java/com/restaurant/entity").glob("*.java")}
        self.assertEqual(names, entities)
        self.assertNotRegex(text, r"\b(?:isAvailable|isActive|getSubtotal|calculateTotal)\(")
        self.assertIn('Order "1" -- "0..1" Payment', text)
        self.assertIn('RestaurantTable "0..1" -- "0..*" Order', text)

    def test_correct_optional_and_mandatory_usecase_relationship_directions(self):
        customer = (DOCS / "bao-cao/uml/usecase-khach-hang.puml").read_text()
        staff = (DOCS / "bao-cao/uml/usecase-nhan-vien.puml").read_text()
        overall = (DOCS / "bao-cao/uml/usecase-tong-quat.puml").read_text()
        self.assertIn('Cashier ..> Receipt : <<include>>', staff)
        self.assertIn('Cashier ..> Record : <<include>>', staff)
        self.assertIn('Voucher .left.> Cashier : <<extend>>', staff)
        self.assertIn('Search .left.> Menu : <<extend>>', customer)
        self.assertIn('AD --|> NV', overall)
        self.assertNotIn('Request ..> ShowQR', customer)


if __name__ == "__main__":
    unittest.main()