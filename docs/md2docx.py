#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Dựng DOCX từ nguồn báo cáo; hỗ trợ bảng, ảnh, ngắt trang và trang ngang.

Cài phụ thuộc: python -m pip install -r docs/requirements.txt
Chạy: python docs/md2docx.py
Các nguồn ảnh thiếu/hỏng làm quá trình dựng thất bại, không ghi đè DOCX cũ.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import os
from pathlib import Path
import re
import struct
import tempfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

BASE = Path(__file__).resolve().parent
SRC = BASE / "BAO_CAO_CHUONG3.md"
OUT = BASE / "BAO_CAO_CHUONG3.docx"
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\n]+?\*)")
IMG_RE = re.compile(r"^!\[(.*?)\]\((.+?)\)\s*$")
TABLE_SEPARATOR_RE = re.compile(r"^:?-{3,}:?$")


def parse_inline(text: str) -> list[tuple[str, str]]:
    result = []
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            result.append(("b", part[2:-2]))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            result.append(("code", part[1:-1]))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            result.append(("i", part[1:-1]))
        else:
            result.append(("n", part))
    return result


def add_inline(paragraph, text: str) -> None:
    for kind, value in parse_inline(text):
        segments = re.split(r"<br\s*/?>", value)
        for index, segment in enumerate(segments):
            run = paragraph.add_run(segment)
            if index:
                # Line break must precede the next segment, inside a run.
                br = OxmlElement("w:br")
                run._r.insert(0 if run._r.rPr is None else 1, br)
            if kind == "b":
                run.bold = True
            elif kind == "i":
                run.italic = True
            elif kind == "code":
                run.font.name = "DejaVu Sans Mono"
                run.font.size = Pt(10)


def set_font(style, size: float, bold: bool = False) -> None:
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{key}"), "Times New Roman")
    language = OxmlElement("w:lang")
    language.set(qn("w:val"), "vi-VN")
    rpr.append(language)


def configure_section(section, landscape: bool = False) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Mm(297 if landscape else 210)
    section.page_height = Mm(210 if landscape else 297)
    section.left_margin = Mm(25)
    section.right_margin = Mm(20)
    section.top_margin = Mm(15 if landscape else 20)
    section.bottom_margin = Mm(15 if landscape else 20)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)


def configure_document(document) -> None:
    configure_section(document.sections[0])
    styles = document.styles
    set_font(styles["Normal"], 13)
    normal = styles["Normal"].paragraph_format
    normal.line_spacing = 1.25
    normal.space_after = Pt(6)
    normal.widow_control = True
    for level, size in ((1, 16), (2, 14), (3, 13), (4, 13)):
        style = styles[f"Heading {level}"]
        set_font(style, size, True)
        fmt = style.paragraph_format
        fmt.keep_with_next = True
        fmt.keep_together = True
        fmt.line_spacing = 1.1
        fmt.space_before = Pt(8 if level > 1 else 0)
        fmt.space_after = Pt(6)
        if level == 1:
            fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(styles["Caption"], 11)
    styles["Caption"].font.italic = True
    caption = styles["Caption"].paragraph_format
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.space_before = Pt(4)
    caption.space_after = Pt(6)
    caption.line_spacing = 1.0
    caption.keep_together = True
    caption.keep_with_next = False
    table_style = styles.add_style("Report Table Text", WD_STYLE_TYPE.PARAGRAPH)
    table_style.base_style = styles["Normal"]
    set_font(table_style, 12)
    table_style.paragraph_format.line_spacing = 1.1
    table_style.paragraph_format.space_after = Pt(3)
    table_style.paragraph_format.space_before = Pt(3)
    table_style.paragraph_format.keep_together = True

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_after = Pt(0)
    text = footer.add_run("Trang ")
    text.font.size = Pt(10)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = "Chương 3 – Phân tích và thiết kế hệ thống quản lý nhà hàng"
    document.core_properties.author = "Restaurant Management System"
    document.core_properties.subject = "Mô hình UML và thiết kế hệ thống"
    document.core_properties.modified = datetime.now(timezone.utc)


def png_size(data: bytes) -> tuple[int, int]:
    if len(data) < 24 or data[:8] != b"\x89PNG\r\n\x1a\n" or data[12:16] != b"IHDR":
        raise ValueError("Ảnh không phải PNG hợp lệ")
    width, height = struct.unpack(">II", data[16:24])
    if not width or not height:
        raise ValueError("Kích thước ảnh PNG phải lớn hơn 0")
    return width, height


def add_image(document, base: Path, rel_path: str, caption: str) -> None:
    path = (base / rel_path).resolve()
    width_px, height_px = png_size(path.read_bytes())
    section = document.sections[-1]
    max_width = section.page_width - section.left_margin - section.right_margin
    max_height = Inches(5.5 if section.orientation == WD_ORIENT.LANDSCAPE else 7.6)
    scale = min(max_width / width_px, max_height / height_px)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = bool(caption)
    picture = paragraph.add_run().add_picture(str(path), width=int(width_px * scale),
                                              height=int(height_px * scale))
    picture._inline.docPr.set("descr", caption)
    if caption:
        document.add_paragraph(caption, style="Caption")


def split_row(line: str) -> list[str]:
    text = line.strip().strip("|")
    return [cell.strip() for cell in text.split("|")]


def is_table_separator(line: str) -> bool:
    return all(TABLE_SEPARATOR_RE.fullmatch(cell) for cell in split_row(line))


def align_of(separator: str):
    if separator.startswith(":") and separator.endswith(":"):
        return WD_ALIGN_PARAGRAPH.CENTER
    if separator.endswith(":"):
        return WD_ALIGN_PARAGRAPH.RIGHT
    return WD_ALIGN_PARAGRAPH.LEFT


def add_table(document, headers: list[str], separators: list[str], rows: list[list[str]]) -> None:
    if len(headers) != len(separators) or any(len(row) != len(headers) for row in rows):
        raise ValueError("Số ô của bảng Markdown không khớp số cột")
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    section = document.sections[-1]
    available = section.page_width - section.left_margin - section.right_margin
    ratios = {2: (0.21, 0.79), 3: (0.22, 0.47, 0.31)}.get(len(headers))
    ratios = ratios or tuple(1 / len(headers) for _ in headers)
    widths = [int(available * ratio) for ratio in ratios]
    for column, width in zip(table.columns, widths):
        column.width = width
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index, values in enumerate([headers, *rows]):
        row = table.rows[0] if index == 0 else table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for column, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = widths[column]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["Report Table Text"]
            paragraph.alignment = align_of(separators[column])
            add_inline(paragraph, text)
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EFEFEF")
                cell._tc.get_or_add_tcPr().append(shade)
                for run in paragraph.runs:
                    run.bold = True
    # A paragraph after a table is required for predictable Word pagination.
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(2)
    spacer.add_run().font.size = Pt(2)


def build_document(markdown: str, base: Path = BASE):
    document = Document()
    configure_document(document)
    lines = markdown.splitlines()
    index = 0
    has_content = False
    pending_break = False
    pending_orientation = None
    while index < len(lines):
        text = lines[index].strip()
        index += 1
        if not text or re.fullmatch(r"-{3,}", text):
            continue
        if text in ("<!-- portrait -->", "<!-- landscape -->"):
            pending_orientation = text == "<!-- landscape -->"
            pending_break = True
            continue
        if text == "<!-- pagebreak -->":
            pending_break = True
            continue
        if text.startswith("<!--"):
            continue
        # Apply consecutive directives together: no empty first/trailing page,
        # and no double page break when an orientation change follows a break.
        current = document.sections[-1].orientation == WD_ORIENT.LANDSCAPE
        break_before = False
        if pending_orientation is not None and pending_orientation != current:
            section = (document.add_section(WD_SECTION_START.NEW_PAGE)
                       if has_content else document.sections[0])
            configure_section(section, pending_orientation)
        elif pending_break and has_content:
            break_before = True
        pending_break = False
        pending_orientation = None
        # New content is inserted before the final body-level sectPr.
        first_block = len(document.element.body) - 1
        image = IMG_RE.fullmatch(text)
        if image:
            add_image(document, base, image.group(2), image.group(1))
        elif (text.startswith("|") and index < len(lines)
              and is_table_separator(lines[index])):
            headers, separators = split_row(text), split_row(lines[index])
            rows = []
            index += 1
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(split_row(lines[index]))
                index += 1
            add_table(document, headers, separators, rows)
        elif text.startswith("```"):
            while index < len(lines) and not lines[index].strip().startswith("```"):
                paragraph = document.add_paragraph()
                run = paragraph.add_run(lines[index])
                run.font.name = "DejaVu Sans Mono"
                run.font.size = Pt(10)
                index += 1
            index += 1
        else:
            heading = re.fullmatch(r"(#{1,4})\s+(.+)", text)
            if heading:
                paragraph = document.add_paragraph(style=f"Heading {len(heading.group(1))}")
                add_inline(paragraph, heading.group(2))
            elif text.startswith(">"):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Mm(5)
                add_inline(paragraph, text[1:].strip())
            elif re.match(r"^[-*]\s+", text):
                paragraph = document.add_paragraph(style="List Bullet")
                add_inline(paragraph, text[2:])
            else:
                paragraph = document.add_paragraph()
                add_inline(paragraph, text)
        if break_before:
            block = document.element.body[first_block]
            if block.tag == qn("w:tbl"):
                block = block.xpath(".//w:p")[0]
            if block.tag == qn("w:p"):
                # Break BEFORE content, not an empty paragraph that can itself
                # overflow and produce a blank page at the end of a full page.
                Paragraph(block, document).paragraph_format.page_break_before = True
        has_content = True
    return document


def convert(source: Path = SRC, output: Path = OUT) -> None:
    source, output = Path(source), Path(output)
    document = build_document(source.read_text(encoding="utf-8"), source.parent)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".docx", delete=False) as temporary:
        temporary_name = temporary.name
    try:
        document.save(temporary_name)
        os.replace(temporary_name, output)
    finally:
        if os.path.exists(temporary_name):
            os.unlink(temporary_name)
    print(f"Đã cập nhật {output} | {len(document.inline_shapes)} ảnh | {len(document.tables)} bảng")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SRC)
    parser.add_argument("--output", type=Path, default=OUT)
    arguments = parser.parse_args()
    convert(arguments.source, arguments.output)
